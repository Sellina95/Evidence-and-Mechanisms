"""Offline, deterministic CASE-002 research build; see design.md. No production imports."""
from pathlib import Path
import hashlib, io, json, re, zipfile, platform
from datetime import datetime, timezone
import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
RAW = ROOT / 'raw'
ASOF = pd.Timestamp('2026-10-05')

def write_csv(df, name):
    df.to_csv(ROOT / name, index=False, float_format='%.10f')

def inputs():
    price = pd.read_excel(RAW/'ypdgdpQvQd.xlsx', index_col=0)
    price.index = pd.PeriodIndex([x.replace(':','') for x in price.index], freq='Q')
    gap = pd.read_excel(RAW/'gap.xlsx', index_col=0)
    gap.index = pd.PeriodIndex([x[:4]+'Q'+str(int(x[-2:])) for x in gap.index], freq='Q')
    dates = pd.read_excel(RAW/'publication_dates.xlsx')
    dates['Greenbook Publication Date'] = pd.to_datetime(dates['Greenbook Publication Date'])
    assert not price.index.has_duplicates and not gap.index.has_duplicates
    return price, gap, dates

def select_signal(q, price, gap, dates):
    cutoff=q.end_time.normalize()
    column=f'YPDGDP{q.year%100:02d}Q{q.quarter}'
    base={'quarter':str(q),'signal_date':str(cutoff.date())}
    if column not in price:
        return base | {'status':'missing_YPDGDP_vintage'}
    availability_bound=pd.Timestamp(q.year, 3*q.quarter-1, 1)+pd.offsets.MonthEnd(0)
    assert availability_bound <= cutoff
    observations=price[column].dropna()
    observations=observations[observations.index<q]
    if observations.empty:
        return base | {'status':'missing_price_observation'}
    s=observations.index[-1]
    a,b=price.at[s,column],price.at[s-4,column]
    if not (pd.notna(a) and pd.notna(b) and a>0 and b>0):
        return base | {'status':'missing_price_four_quarter_pair'}
    eligible=[]
    for c in gap:
        column_date=pd.to_datetime(c.split('_')[1],format='%y%m%d')
        if column_date.year<2015: continue
        candidates=dates.loc[(dates['Greenbook Publication Date']-column_date).abs()<=pd.Timedelta(days=7)]
        assert len(candidates)==1, f'Ambiguous staff-date mapping {c}'
        publication_date=candidates.iloc[0]['Greenbook Publication Date']
        available=max(column_date,publication_date)
        if available<=cutoff: eligible.append((available,c,column_date,publication_date,candidates.iloc[0]['FOMC Meeting']))
    eligible.sort()
    if not eligible or eligible[-1][0].to_period('Q')!=q:
        return base | {'status':'missing_staff_vintage_in_signal_quarter'}
    gd,gc,column_date,publication_date,meeting=eligible[-1]
    g=gap.at[s,gc] if s in gap.index else np.nan
    if pd.isna(g):
        return base | {'status':'missing_gap_at_matched_economic_quarter'}
    pi=100*(a/b-1)
    pc=get_column_letter(price.columns.get_loc(column)+2)
    gg=get_column_letter(gap.columns.get_loc(gc)+2)
    return base | dict(status='ok', information_set='Fed_staff_PIT_not_public_market_PIT',
        economic_quarter=str(s),economic_lag_quarters=q.ordinal-s.ordinal,price_vintage=column,
        price_available_by=str(availability_bound.date()),price_availability_precision='conservative_month_end_bound',
        price_level=a,price_level_four_quarters_earlier=b,
        price_cell=f'YPDGDP!{pc}{price.index.get_loc(s)+2}',
        price_lag_cell=f'YPDGDP!{pc}{price.index.get_loc(s-4)+2}',
        inflation_pct=pi,staff_vintage=gc,staff_available_date=str(gd.date()),
        staff_column_date=str(column_date.date()),staff_publication_date=str(publication_date.date()),
        staff_date_mapping='exact' if column_date==publication_date else 'unique_within_7_days_later_bound_not_exact_verification',
        staff_meeting=meeting,gap_cell=f'gap!{gg}{gap.index.get_loc(s)+2}',
        output_gap_pct=g,public_gap_release_date=None,
        taylor_pct=2+pi+.5*(pi-2)+.5*g)

def market_data():
    # FRED serves a ZIP for mixed daily / daily-seven-day frequencies.
    with zipfile.ZipFile(RAW/'rates.csv') as z:
        frames=[pd.read_csv(io.BytesIO(z.read(n)),parse_dates=['observation_date']).set_index('observation_date') for n in ['daily,_7-day.csv','daily.csv']]
        (ROOT/'fred_download_notes.txt').write_bytes(z.read('README.txt').rstrip()+b'\n')
    daily=pd.concat(frames,axis=1).sort_index()
    acm=pd.read_excel(RAW/'ACMTermPremium.xls',sheet_name='ACM Daily')
    acm.index=pd.to_datetime(acm.pop('DATE'),format='%d-%b-%Y')
    daily=daily.join(acm[['ACMTP10','ACMY10','ACMRNY10']],how='outer').sort_index()
    daily=daily.loc[daily.index<ASOF].apply(pd.to_numeric,errors='raise')
    assert not daily.index.has_duplicates
    rows=[]
    coverage=[]
    for col in daily:
        non=daily[col].dropna()
        coverage.append(dict(series=col,first_observation=str(non.index.min().date()),last_observation=str(non.index.max().date()),n=len(non),vintage_status='current_download_not_historical_release_archive'))
    write_csv(pd.DataFrame(coverage),'source_coverage.csv')
    for q,d in daily.groupby(daily.index.to_period('Q')):
        if q.end_time.normalize()>=ASOF:
            continue
        row={'quarter':str(q)}
        for c in daily:
            s=d[c].dropna()
            row[c]=s.mean()
            row[c+'_n']=len(s)
            row[c+'_last_date']=str(s.index.max().date()) if len(s) else None
        row['DFF_complete_calendar_quarter']=row['DFF_n']==(q.end_time.normalize()-q.start_time).days+1
        rows.append(row)
    return daily,pd.DataFrame(rows).set_index('quarter')

def metrics(df, label, horizon):
    out=[]
    for c in ['DFF','DGS2','DGS10']:
        target=c if horizon==0 else c+'_next'
        d=df[['taylor_pct',target]].dropna()
        residual=d.taylor_pct-d[target]
        changes=d.diff().dropna()
        row=dict(comparison=label,target=c,n=len(d),mean_taylor_minus_rate_pp=residual.mean(),mae_pp=residual.abs().mean(),rmse_pp=np.sqrt((residual**2).mean()),level_correlation=d.taylor_pct.corr(d[target]),change_correlation=changes.taylor_pct.corr(changes[target]),n_changes=len(changes))
        if horizon:
            e=df[c]-df[target]
            row['prior_realized_mean_no_change_mae_pp']=e.abs().mean()
            row['prior_realized_mean_no_change_rmse_pp']=np.sqrt((e**2).mean())
        out.append(row)
    return out

def main():
    price,gap,dates=inputs()
    quarters=pd.period_range('1996Q1','2026Q3',freq='Q')
    ledger=pd.DataFrame([select_signal(q,price,gap,dates) for q in quarters])
    write_csv(ledger,'coverage_ledger.csv')
    signals=ledger[ledger.status=='ok'].copy()
    assert len(signals)==23 and signals.quarter.iloc[0]=='2015Q2' and signals.quarter.iloc[-1]=='2020Q4'
    assert (pd.to_datetime(signals.staff_available_date)<=pd.to_datetime(signals.signal_date)).all()
    assert (pd.to_datetime(signals.price_available_by)<=pd.to_datetime(signals.signal_date)).all()
    assert all(pd.Period(s,'Q')<pd.Period(q,'Q') for s,q in zip(signals.economic_quarter,signals.quarter))
    # Leakage regression: corrupt every future vintage. Historical result must not change.
    p2,g2=price.copy(),gap.copy()
    qtest=pd.Period('2018Q2','Q')
    for c in p2:
        if pd.Period('20'+c[6:8]+'Q'+c[-1],'Q')>qtest: p2[c]=999999.
    for c in g2:
        if pd.to_datetime(c.split('_')[1],format='%y%m%d')>qtest.end_time: g2[c]=999999.
    assert select_signal(qtest,price,gap,dates)==select_signal(qtest,p2,g2,dates)
    # Missing contemporaneous gap must remain missing, never backfilled from a future column.
    chosen=select_signal(qtest,price,gap,dates)
    g3=gap.copy(); g3.at[pd.Period(chosen['economic_quarter'],'Q'),chosen['staff_vintage']]=np.nan
    assert select_signal(qtest,price,g3,dates)['status']=='missing_gap_at_matched_economic_quarter'
    daily,market=market_data()
    write_csv(market.reset_index(),'quarterly_market_diagnostics.csv')
    same=signals.merge(market.reset_index(),on='quarter',validate='one_to_one')
    assert same.DFF_complete_calendar_quarter.all()
    assert same[['DGS2_n','DGS10_n','DGS30_n','DFII10_n','T10YIE_n','ACMTP10_n']].ge(55).all().all()
    same['taylor_minus_effr_pp']=same.taylor_pct-same.DFF
    write_csv(same,'quarterly_taylor_comparison.csv')
    following=signals.copy()
    following['outcome_quarter']=[str(pd.Period(q,'Q')+1) for q in following.quarter]
    following=following.merge(market[['DFF','DGS2','DGS10']].reset_index(),on='quarter',validate='one_to_one')
    following=following.merge(market[['DFF','DGS2','DGS10']].add_suffix('_next').rename_axis('outcome_quarter').reset_index(),on='outcome_quarter',validate='one_to_one')
    assert all(pd.Period(q,'Q').start_time>pd.Timestamp(s) for q,s in zip(following.outcome_quarter,following.signal_date))
    write_csv(following,'following_quarter_comparison.csv')
    stats=pd.DataFrame(metrics(same,'same_quarter_descriptive',0)+metrics(following,'following_quarter_realization',1))
    write_csv(stats,'comparison_metrics.csv')
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig,axes=plt.subplots(2,1,figsize=(11,7.5),sharex=True,layout='constrained')
    x=pd.PeriodIndex(same.quarter,freq='Q').to_timestamp(how='end')
    colors={'taylor_pct':'#9b3d10','DFF':'#172d48','DGS2':'#007d91','DGS10':'#77549c'}
    labels={'taylor_pct':'Taylor (quarter-end staff PIT)','DFF':'EFFR quarterly mean','DGS2':'2Y quarterly mean','DGS10':'10Y quarterly mean'}
    for c in colors:
        axes[0].plot(x,same[c],label=labels[c],color=colors[c],linewidth=2)
        axes[1].plot(x,100*same[c].diff(),color=colors[c],linewidth=1.6)
    axes[0].set_title('CASE-002 | First frozen Taylor-93-style comparison',loc='left',fontweight='bold',fontsize=15)
    axes[0].set_ylabel('Rate (%)'); axes[1].set_ylabel('Quarterly change (bp)')
    axes[0].legend(loc='upper left',fontsize=9,ncol=2)
    for ax in axes:
        ax.axhline(0,color='#777777',linewidth=.7); ax.grid(axis='y',alpha=.2)
        ax.spines[['top','right']].set_visible(False)
    fig.supxlabel('2015Q2–2020Q4 | Same-quarter comparison is descriptive, not a forecast.\nLagged economic inputs; staff information was not public then. Sources: Philadelphia Fed and FRED.',fontsize=9)
    fig.savefig(ROOT/'first_comparison.png',dpi=160)
    plt.close(fig)
    match=daily.loc['2015-04-01':'2020-12-31',['DGS10','DFII10','T10YIE']].dropna()
    identity=match.DGS10-match.DFII10-match.T10YIE
    # Source vintage coverage and row ranges are different concepts.
    audit=dict(signal_count=len(signals),signal_start=signals.quarter.iloc[0],signal_end=signals.quarter.iloc[-1],
        no_future_vintage_test='passed',missing_input_no_backfill_test='passed',
        date_crosscheck=signals.staff_date_mapping.value_counts().to_dict(),
        DFF_calendar_completeness='passed',market_minimum_observations='55 per quarter passed',
        matched_daily_breakeven_identity_max_abs_pp=float(identity.abs().max()),
        economic_lag_quarters_counts=signals.economic_lag_quarters.value_counts().to_dict(),
        price_first_vintage=price.columns[0],price_last_vintage=price.columns[-1],
        staff_first_vintage=gap.columns[0],staff_last_vintage=gap.columns[-1],
        coverage_status_counts=ledger.status.value_counts().to_dict(),
        python=platform.python_version(),pandas=pd.__version__,numpy=np.__version__)
    (ROOT/'validation.json').write_text(json.dumps(audit,indent=2)+'\n')
    manifest=[]
    for name,url in json.loads((ROOT/'sources.json').read_text()).items():
        f=RAW/name
        manifest.append(dict(file='raw/'+name,url=url,bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest(),download_completed_utc=datetime.fromtimestamp(f.stat().st_mtime,timezone.utc).isoformat(),timestamp_basis='local file mtime at first manifest creation'))
    mf=ROOT/'manifest.json'
    if not mf.exists(): mf.write_text(json.dumps(manifest,indent=2)+'\n')
    else:
        recorded=json.loads(mf.read_text())
        assert [(r['file'],r['sha256']) for r in recorded]==[(r['file'],r['sha256']) for r in manifest], 'Raw input changed; create explicit new snapshot.'
    print(same[['quarter','economic_quarter','inflation_pct','output_gap_pct','taylor_pct','DFF','DGS2','DGS10']].to_string(index=False))
    print(stats.to_string(index=False))
    print(json.dumps(audit,indent=2))

if __name__=='__main__': main()
