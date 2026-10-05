"""Independently reconcile exported results to named raw Excel cells and FRED rows."""
from pathlib import Path
from decimal import Decimal
import csv, io, json, zipfile
from openpyxl import load_workbook

root=Path(__file__).resolve().parent
price=load_workbook(root/'raw/ypdgdpQvQd.xlsx',data_only=True)['YPDGDP']
gap=load_workbook(root/'raw/gap.xlsx',data_only=True)['gap']
rows=list(csv.DictReader((root/'quarterly_taylor_comparison.csv').open()))
with zipfile.ZipFile(root/'raw/rates.csv') as z:
    fred=list(csv.DictReader(io.StringIO(z.read('daily,_7-day.csv').decode())))
for row in rows:
    a=Decimal(str(price[row['price_cell'].split('!')[1]].value))
    b=Decimal(str(price[row['price_lag_cell'].split('!')[1]].value))
    g=Decimal(str(gap[row['gap_cell'].split('!')[1]].value))
    pi=100*(a/b-1)
    rate=Decimal(2)+pi+Decimal('.5')*(pi-2)+Decimal('.5')*g
    assert abs(rate-Decimal(row['taylor_pct']))<Decimal('0.000000001')
    year=int(row['quarter'][:4]); quarter=int(row['quarter'][-1])
    values=[Decimal(r['DFF']) for r in fred if int(r['observation_date'][:4])==year and (int(r['observation_date'][5:7])-1)//3+1==quarter and r['DFF']]
    assert abs(sum(values)/len(values)-Decimal(row['DFF']))<Decimal('0.000000001')
result={'raw_cell_formula_reconciliation_rows':len(rows),'independent_calendar_DFF_mean_rows':len(rows),'absolute_tolerance_pp':1e-9,'status':'passed'}
(root/'independent_validation.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result))
